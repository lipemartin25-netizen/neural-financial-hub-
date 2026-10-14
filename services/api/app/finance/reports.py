import io
import csv
from datetime import datetime
from typing import List, Dict, Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False


def generate_csv_report(
    transactions: List[Dict[str, Any]],
    accounts_map: Dict[str, str],
    categories_map: Dict[str, str]
) -> str:
    """Gera relatório financeiro em formato CSV padronizado com resumo."""
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";", quoting=csv.QUOTE_MINIMAL)

    # Metadados
    writer.writerow(["# NEURAFIN HUB - RELATORIO FINANCEIRO CONSOLIDADO"])
    writer.writerow([f"# Gerado em: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
    writer.writerow(["# Modo: Financas Manuais Privadas (Sem Open Finance / Sem Pluggy)"])
    writer.writerow([])

    # Cabeçalho
    writer.writerow(["Data", "Tipo", "Descricao", "Categoria", "Conta", "Valor (R$)", "Confirmado"])

    total_income = 0.0
    total_expense = 0.0

    for tx in transactions:
        tx_type = tx.get("type", "expense")
        amount = float(tx.get("amount", 0.0))
        if tx_type == "income":
            total_income += amount
        elif tx_type == "expense":
            total_expense += amount

        cat_name = categories_map.get(str(tx.get("category_id")), "Geral")
        acc_name = accounts_map.get(str(tx.get("account_id")), "Conta Principal")
        confirmed = "Sim" if tx.get("is_confirmed_by_user", True) else "Nao"

        writer.writerow([
            tx.get("date", ""),
            "Receita" if tx_type == "income" else ("Despesa" if tx_type == "expense" else "Transferencia"),
            tx.get("description", ""),
            cat_name,
            acc_name,
            f"{amount:.2f}",
            confirmed
        ])

    net_balance = total_income - total_expense
    writer.writerow([])
    writer.writerow(["# RESUMO CONSOLIDADO"])
    writer.writerow(["Total Receitas (R$)", f"{total_income:.2f}"])
    writer.writerow(["Total Despesas (R$)", f"{total_expense:.2f}"])
    writer.writerow(["Saldo Liquido (R$)", f"{net_balance:.2f}"])

    return output.getvalue()


def generate_excel_report(
    transactions: List[Dict[str, Any]],
    accounts_map: Dict[str, str],
    categories_map: Dict[str, str]
) -> bytes:
    """Gera relatório financeiro em formato XLSX (Excel) com fórmulas reais verificadas."""
    if not OPENPYXL_AVAILABLE:
        raise RuntimeError("openpyxl não está instalado no ambiente.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Extrato Financeiro"

    # Estilos
    header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    bold_font = Font(name="Calibri", size=11, bold=True)
    border_thin = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    # Título institucional
    ws["A1"] = "NEURAFIN HUB — RELATÓRIO FINANCEIRO CONSOLIDADO"
    ws["A1"].font = title_font
    ws["A2"] = f"Exportado em: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Modo: Privado e Manual (Sem Open Finance)"
    ws["A2"].font = subtitle_font

    # Cabeçalho da tabela
    headers = ["Data", "Tipo", "Descrição", "Categoria", "Conta", "Valor (R$)", "Confirmado"]
    start_row = 4
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 7] else ("right" if col_idx == 6 else "left"))
        cell.border = border_thin

    # Linhas de dados
    curr_row = start_row + 1
    for tx in transactions:
        tx_type = tx.get("type", "expense")
        amount = float(tx.get("amount", 0.0))
        cat_name = categories_map.get(str(tx.get("category_id")), "Geral")
        acc_name = accounts_map.get(str(tx.get("account_id")), "Conta Principal")
        confirmed = "Sim" if tx.get("is_confirmed_by_user", True) else "Não"

        ws.cell(row=curr_row, column=1, value=tx.get("date", "")).alignment = Alignment(horizontal="center")
        ws.cell(row=curr_row, column=2, value="Receita" if tx_type == "income" else ("Despesa" if tx_type == "expense" else "Transferência")).alignment = Alignment(horizontal="center")
        ws.cell(row=curr_row, column=3, value=tx.get("description", ""))
        ws.cell(row=curr_row, column=4, value=cat_name)
        ws.cell(row=curr_row, column=5, value=acc_name)

        amount_cell = ws.cell(row=curr_row, column=6, value=amount)
        amount_cell.number_format = '"R$" #,##0.00'
        amount_cell.alignment = Alignment(horizontal="right")

        ws.cell(row=curr_row, column=7, value=confirmed).alignment = Alignment(horizontal="center")

        for c in range(1, 8):
            ws.cell(row=curr_row, column=c).border = border_thin

        curr_row += 1

    last_data_row = curr_row - 1

    # Seção de Resumo com Fórmulas Excel Reais
    curr_row += 1
    summary_start = curr_row

    ws.cell(row=curr_row, column=4, value="Total Receitas").font = bold_font
    income_formula_cell = ws.cell(
        row=curr_row,
        column=6,
        value=f'=SUMIF(B{start_row + 1}:B{last_data_row}, "Receita", F{start_row + 1}:F{last_data_row})' if last_data_row >= start_row + 1 else 0.0
    )
    income_formula_cell.font = bold_font
    income_formula_cell.number_format = '"R$" #,##0.00'

    curr_row += 1
    ws.cell(row=curr_row, column=4, value="Total Despesas").font = bold_font
    expense_formula_cell = ws.cell(
        row=curr_row,
        column=6,
        value=f'=SUMIF(B{start_row + 1}:B{last_data_row}, "Despesa", F{start_row + 1}:F{last_data_row})' if last_data_row >= start_row + 1 else 0.0
    )
    expense_formula_cell.font = bold_font
    expense_formula_cell.number_format = '"R$" #,##0.00'

    curr_row += 1
    ws.cell(row=curr_row, column=4, value="Saldo Líquido").font = bold_font
    net_formula_cell = ws.cell(
        row=curr_row,
        column=6,
        value=f'=F{summary_start}-F{summary_start + 1}'
    )
    net_formula_cell.font = bold_font
    net_formula_cell.number_format = '"R$" #,##0.00'

    # Auto-ajuste de largura de colunas
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_pdf_report(
    transactions: List[Dict[str, Any]],
    accounts_map: Dict[str, str],
    categories_map: Dict[str, str]
) -> bytes:
    """Gera relatório financeiro institucional em PDF formatado."""
    if not REPORTLAB_AVAILABLE:
        raise RuntimeError("reportlab não está instalado no ambiente.")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#0f172a")
    )
    subtitle_style = ParagraphStyle(
        "ReportSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#64748b")
    )
    cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#1e293b")
    )
    header_cell_style = ParagraphStyle(
        "HeaderCell",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        textColor=colors.white
    )

    story = []

    # Título e Subtítulo
    story.append(Paragraph("NEURAFIN HUB — RELATÓRIO FINANCEIRO CONSOLIDADO", title_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph(
        f"Emissão: {datetime.now().strftime('%d/%m/%Y às %H:%M')} | Governança: Cofre Local & Lançamentos Manuais (Sem Agregadores de Open Finance)",
        subtitle_style
    ))
    story.append(Spacer(1, 14))

    # Totais calculados
    total_income = sum(float(t.get("amount", 0.0)) for t in transactions if t.get("type") == "income")
    total_expense = sum(float(t.get("amount", 0.0)) for t in transactions if t.get("type") == "expense")
    net_balance = total_income - total_expense

    # Tabela de Resumo (Cards)
    summary_data = [
        [
            Paragraph("<b>Total Receitas</b>", cell_style),
            Paragraph(f"<font color='#10b981'><b>R$ {total_income:,.2f}</b></font>", cell_style),
            Paragraph("<b>Total Despesas</b>", cell_style),
            Paragraph(f"<font color='#ef4444'><b>R$ {total_expense:,.2f}</b></font>", cell_style),
            Paragraph("<b>Saldo Líquido</b>", cell_style),
            Paragraph(f"<b>R$ {net_balance:,.2f}</b>", cell_style)
        ]
    ]
    summary_table = Table(summary_data, colWidths=[90, 85, 90, 85, 85, 85])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#e2e8f0")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 14))

    # Tabela de Lançamentos
    table_rows = [
        [
            Paragraph("Data", header_cell_style),
            Paragraph("Tipo", header_cell_style),
            Paragraph("Descrição", header_cell_style),
            Paragraph("Categoria", header_cell_style),
            Paragraph("Conta", header_cell_style),
            Paragraph("Valor (R$)", header_cell_style),
            Paragraph("Conf.", header_cell_style)
        ]
    ]

    for tx in transactions:
        tx_type = tx.get("type", "expense")
        amount = float(tx.get("amount", 0.0))
        cat_name = categories_map.get(str(tx.get("category_id")), "Geral")
        acc_name = accounts_map.get(str(tx.get("account_id")), "Principal")
        confirmed = "✓" if tx.get("is_confirmed_by_user", True) else "—"

        type_label = "Receita" if tx_type == "income" else ("Despesa" if tx_type == "expense" else "Transf.")
        type_color = "#10b981" if tx_type == "income" else ("#ef4444" if tx_type == "expense" else "#64748b")

        table_rows.append([
            Paragraph(tx.get("date", ""), cell_style),
            Paragraph(f"<font color='{type_color}'>{type_label}</font>", cell_style),
            Paragraph(tx.get("description", "")[:35], cell_style),
            Paragraph(cat_name[:20], cell_style),
            Paragraph(acc_name[:18], cell_style),
            Paragraph(f"{amount:,.2f}", cell_style),
            Paragraph(confirmed, cell_style)
        ])

    tx_table = Table(table_rows, colWidths=[55, 45, 175, 75, 75, 65, 30])
    tx_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("ALIGN", (5, 0), (5, -1), "RIGHT"),
        ("ALIGN", (6, 0), (6, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(tx_table)

    doc.build(story)
    return buf.getvalue()
